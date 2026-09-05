import io
from datetime import datetime, timezone, timedelta

import pytest
from openpyxl import load_workbook

from catalog import Category, CategoryPart
from excel import (
    HEADERS,
    PriceRow,
    build_caption,
    build_filename,
    build_workbook,
    collect_rows,
    rows_hash,
)
from markup import parse_tiers

TIERS = parse_tiers([{"up_to": 80000, "add": 5000}, {"up_to": None, "add": 13000}])


def make_category(title: str, lines: list[str]) -> Category:
    part = CategoryPart(message_id=1, part_no=1, declared_total=1, lines=tuple(lines))
    return Category(title=title, key=title.casefold(), parts=[part])


def test_collect_rows_applies_markup_and_skips_non_products():
    category = make_category("iPhone 17 Pro", [
        "17 Pro 256GB Blue 🇯🇵 (E-Sim) - 95900",
        "",
        "какой-то текст без цены",
        "17 Pro 512GB Blue 🇰🇷 (Sim + E-Sim) - 128000",
    ])
    rows = collect_rows([category], TIERS)
    assert [(r.name, r.price) for r in rows] == [
        ("17 Pro 256GB Blue 🇯🇵 (E-Sim)", 108900),
        ("17 Pro 512GB Blue 🇰🇷 (Sim + E-Sim)", 141000),
    ]
    assert all(r.category == "iPhone 17 Pro" for r in rows)


def test_collect_rows_keeps_category_order():
    rows = collect_rows(
        [make_category("A", ["x - 10000"]), make_category("B", ["y - 20000"])], TIERS,
    )
    assert [r.category for r in rows] == ["A", "B"]


def test_rows_hash_is_stable_and_sensitive():
    a = [PriceRow("C", "N", 100)]
    assert rows_hash(a) == rows_hash([PriceRow("C", "N", 100)])
    assert rows_hash(a) != rows_hash([PriceRow("C", "N", 101)])
    assert rows_hash(a) != rows_hash([PriceRow("C", "M", 100)])


def test_build_workbook_produces_readable_sheet():
    rows = [PriceRow("iPhone 17 Pro", "17 Pro 256GB Blue", 108900),
            PriceRow("AirPods", "AirPods 4 (Type-C)", 11600)]
    book = load_workbook(io.BytesIO(build_workbook(rows)))
    sheet = book.active
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == 3
    assert sheet["C2"].value == 108900
    assert isinstance(sheet["C2"].value, int)
    assert sheet["B3"].value == "AirPods 4 (Type-C)"


def test_build_workbook_sets_filter_and_freeze():
    book = load_workbook(io.BytesIO(build_workbook([PriceRow("C", "N", 100)])))
    sheet = book.active
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C2"


def test_build_workbook_handles_empty_rows():
    book = load_workbook(io.BytesIO(build_workbook([])))
    assert book.active.max_row == 1


def test_build_filename_substitutes_date():
    moment = datetime(2026, 9, 5, tzinfo=timezone(timedelta(hours=3)))
    assert build_filename("Прайс_{date}.xlsx", moment) == "Прайс_2026-09-05.xlsx"


def test_build_caption_reports_count_and_time():
    moment = datetime(2026, 9, 5, 18, 30, tzinfo=timezone(timedelta(hours=3)))
    caption = build_caption([PriceRow("C", "N", 100)], moment)
    assert "Позиций: 1" in caption
    assert "05.09.2026 18:30" in caption
