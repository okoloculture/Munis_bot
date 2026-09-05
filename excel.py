"""Выгрузка прайса в .xlsx."""

from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from catalog import Category
from markup import MarkupTier, apply_markup
from renderer import parse_position

MSK = timezone(timedelta(hours=3))
SHEET_TITLE = "Прайс"
HEADERS = ("Категория", "Наименование", "Цена, ₽")
COLUMN_WIDTHS = (26, 58, 14)
PRICE_NUMBER_FORMAT = "# ##0"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass(frozen=True)
class PriceRow:
    """Одна позиция прайса с уже применённой наценкой."""

    category: str
    name: str
    price: int


def collect_rows(
    categories: Sequence[Category],
    tiers: Sequence[MarkupTier],
    round_to: int = 0,
) -> list[PriceRow]:
    """Развернуть категории в плоский список позиций с наценкой."""
    rows: list[PriceRow] = []
    for category in categories:
        for line in category.lines:
            parsed = parse_position(line)
            if parsed is None:
                continue
            name, price = parsed
            rows.append(
                PriceRow(
                    category=category.title,
                    name=name,
                    price=apply_markup(price, tiers, round_to),
                )
            )
    return rows


def rows_hash(rows: Iterable[PriceRow]) -> str:
    """Хеш содержимого выгрузки: по данным, а не по байтам файла.

    Два .xlsx с одинаковыми позициями отличаются байтами (время создания, порядок
    записей в архиве), поэтому сравнивать надо именно строки.
    """
    digest = hashlib.md5()
    for row in rows:
        digest.update(f"{row.category}\x1f{row.name}\x1f{row.price}\x1e".encode("utf-8"))
    return digest.hexdigest()


def build_workbook(rows: Sequence[PriceRow]) -> bytes:
    """Собрать книгу Excel: один лист, шапка с фильтром, цены числами."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET_TITLE

    sheet.append(list(HEADERS))
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append([row.category, row.name, row.price])

    for index, width in enumerate(COLUMN_WIDTHS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    price_column = get_column_letter(len(HEADERS))
    for cell in sheet[price_column][1:]:
        cell.number_format = PRICE_NUMBER_FORMAT

    last_row = sheet.max_row
    sheet.auto_filter.ref = f"A1:{price_column}{last_row}"
    sheet.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def build_filename(template: str, now: datetime | None = None) -> str:
    """Имя файла по шаблону; {date} подставляется датой по Москве."""
    moment = now or datetime.now(MSK)
    return template.format(date=moment.strftime("%Y-%m-%d"))


def build_caption(rows: Sequence[PriceRow], now: datetime | None = None) -> str:
    """Подпись к файлу в канале."""
    moment = now or datetime.now(MSK)
    return (
        f"<b>Прайс-лист</b>\n"
        f"Позиций: {len(rows)}\n"
        f"Обновлено: {moment.strftime('%d.%m.%Y %H:%M')} МСК"
    )
